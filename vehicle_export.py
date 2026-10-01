"""Excel + PDF export of the /vehicles fuel-cost analysis (vehicle_costs.build_page_data()).

Reuses the look of the expense report in export_report so both read as one set.
"""

from datetime import date, datetime
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.charts.linecharts import HorizontalLineChart
from reportlab.graphics.shapes import Drawing, Rect, String
from reportlab.lib.pagesizes import A4
from reportlab.platypus import CondPageBreak, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from export_report import (
    BANGKOK_TZ, C_ACCENT, C_ACCENT_SOFT, C_DOWN, C_INK, C_MUTED, C_SURFACE_ALT, C_UP, C_WARN, CONTENT_W, FONT,
    FONT_BOLD, MARGIN, MONEY, ST, P, PW, _base_table_style, _callout, _f, _fill, _hex, _medium,
    _NumberedCanvas, _nice_max, _put, _thin, _xl_header, _xl_print_setup, _xl_section, _xl_title, fmt_baht,
    fmt_date, header_band, kpi_cards,
)

TITLE = "รายงานต้นทุนรถ"


def _d(s):
    return date.fromisoformat(s)


def _meta(data):
    k = data["kpi"]
    gen = datetime.now(BANGKOK_TZ)
    rng = f"{fmt_date(_d(k['first_date']))} – {fmt_date(_d(k['last_date']))}"
    # Shape expected by export_report's sheet helpers (_xl_title / _xl_print_setup).
    return {"label": "ทุกคัน", "range_label": rng, "generated_at": gen, "scope": None}


def export_filename(data, ext):
    return f"vehicle_costs_{data['kpi']['last_date']}.{ext}"


# Same order as the /vehicles page chart, so a vehicle keeps its colour across web, PDF and Excel.
VEHICLE_COLORS = ["2A78D6", "1BAF7A", "EDA100", "4A3AA7", "E34948", "E87BA4", "EB6834", "008300"]


def _color(i):
    return VEHICLE_COLORS[i % len(VEHICLE_COLORS)]


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def _xl_summary(wb, data, meta):
    ws = wb.active
    ws.title = "สรุป"
    k = data["kpi"]
    widths = {"A": 24, "B": 10, "C": 11, "D": 13, "E": 11, "F": 16, "G": 8, "H": 14, "I": 10}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    last = 9
    _xl_title(ws, meta, TITLE, last)

    row = _xl_section(ws, 4, "ภาพรวม", last)
    d_now, d_base = data["latest_price"].get("diesel"), data["baseline"].get("diesel")
    lines = [
        ("ค่าน้ำมันรวม", k["amount"], MONEY, f"บาท · {k['fills']} ครั้ง · {k['liters']:,.0f} ลิตร"),
        ("เฉลี่ยต่อเดือน (3 เดือนล่าสุด)", k["monthly_avg_recent"], MONEY,
         f"บาท · ~{k['monthly_liters_recent']:,.0f} ลิตร/เดือน (ดีเซล ±1 บาท = ±{k['monthly_liters_recent']:,.0f} บาท/เดือน)"),
        (f"จ่ายเพิ่มเพราะราคาน้ำมัน (ตั้งแต่ {data['baseline_label']})", k["price_effect_since"], MONEY,
         f"บาท เทียบราคาเฉลี่ยก่อนหน้า (ดีเซล {d_base:.2f})" if d_base else "บาท"),
        ("ราคาดีเซลล่าสุด", d_now or "–", "0.00", "บาท/ลิตร"),
    ]
    for label, value, fmt, note in lines:
        _put(ws, row, 1, label, _f(10, color=C_MUTED), align=Alignment(indent=1, vertical="center"))
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        bold_color = C_UP if label.startswith("จ่ายเพิ่ม") else C_INK
        _put(ws, row, 2, value, _f(12, True, bold_color), fmt=fmt, align=Alignment(horizontal="right", vertical="center"))
        ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=last)
        _put(ws, row, 4, note, _f(9, color=C_MUTED), align=Alignment(indent=1, vertical="center"))
        for col in range(1, last + 1):
            ws.cell(row=row, column=col).border = Border(bottom=_thin)
        ws.row_dimensions[row].height = 21
        row += 1
    for v in data["vehicles"]:
        if v["missing_liters"]:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last)
            _put(ws, row, 1, f"{v['name']}: ใบเสร็จขาดประมาณ {v['missing_liters']:,.0f} ลิตร "
                             f"(~{v['missing_liters'] * (d_now or 0):,.0f} บาท ที่ราคาปัจจุบัน)",
                 _f(10, True, C_WARN), _fill("FEF3E2"), align=Alignment(indent=1, vertical="center"))
            ws.row_dimensions[row].height = 22
            row += 1

    row = _xl_section(ws, row + 1, "เทียบรายคัน", last)
    row = _xl_header(ws, row, ["รถ", "น้ำมัน", "กม./ลิตร", "วิ่ง/เดือน (กม.)", "บาท/กม.", "ค่าน้ำมันรวม",
                               "เติม", "เลขไมล์ล่าสุด", "ต้องตรวจ"], "llrrrrrrr")
    for v in data["vehicles"]:
        _put(ws, row, 1, f"{v['name']}  {v['plate']}", _f(10, True), align=Alignment(indent=1))
        _put(ws, row, 2, v["fuel_label"])
        _put(ws, row, 3, v["km_per_l"], fmt="0.0")
        _put(ws, row, 4, v["km_per_month"], fmt="#,##0")
        _put(ws, row, 5, v["baht_per_km"], _f(10, True), fmt="0.00")
        _put(ws, row, 6, v["amount"], fmt=MONEY)
        _put(ws, row, 7, v["fills"], fmt="#,##0")
        _put(ws, row, 8, v["odometer_last"], fmt="#,##0")
        _put(ws, row, 9, v["issues"] or "–", _f(10, bool(v["issues"]), C_WARN if v["issues"] else C_MUTED),
             align=Alignment(horizontal="right"))
        for col in range(1, last + 1):
            ws.cell(row=row, column=col).border = Border(bottom=_thin)
        row += 1
        if v["efficiency_note"]:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last)
            _put(ws, row, 1, v["efficiency_note"], _f(9, color=C_WARN, italic=True), align=Alignment(indent=2, wrap_text=True))
            ws.row_dimensions[row].height = 28
            row += 1

    u = data["usage"]
    row = _xl_section(ws, row + 1, "ราคาขึ้นแล้ว เราใช้น้ำมันเปลี่ยนไหม", last)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last)
    _put(ws, row, 1, usage_headline(u), _f(10, True, C_ACCENT), _fill(C_ACCENT_SOFT),
         align=Alignment(indent=1, vertical="center", wrap_text=True))
    ws.row_dimensions[row].height = 32
    row += 1
    _put(ws, row, 1, f"ก่อน = {u['before_label']} · หลัง = {u['after_label']} (เฉลี่ยต่อเดือน)",
         _f(9, color=C_MUTED, italic=True), align=Alignment(indent=1))
    row += 1
    row = _xl_header(ws, row, ["รถ", "ลิตร/ด. ก่อน", "หลัง", "เปลี่ยน", "ผล", "บาท/ด. เปลี่ยน", "", "จากปริมาณ", "จากราคา"],
                     "lrrrlrrrr")
    signed = "+#,##0;-#,##0;0"
    for r in u["vehicles"]:
        ch = r["liters_change"]
        _put(ws, row, 1, r["name"], _f(10, True), align=Alignment(indent=1))
        _put(ws, row, 2, r["before"]["liters"], fmt="#,##0")
        _put(ws, row, 3, r["after"]["liters"], fmt="#,##0")
        _put(ws, row, 4, ch / 100 if ch is not None else "–",
             _f(10, False, C_UP if ch and ch >= 5 else C_DOWN if ch and ch <= -5 else C_INK), fmt="+0.0%;-0.0%;0.0%")
        _put(ws, row, 5, "ใบเสร็จไม่ครบ" if r["missing_liters"] else r["verdict"],
             _f(10, color=C_MUTED if r["missing_liters"] else C_INK))
        _put(ws, row, 6, r["spend_change"], fmt=signed)
        _put(ws, row, 8, r["volume_effect"], _f(10, color=C_UP if r["volume_effect"] > 0.5 else C_DOWN), fmt=signed)
        _put(ws, row, 9, r["price_effect"], _f(10, color=C_UP if r["price_effect"] > 0.5 else C_DOWN), fmt=signed)
        for col in range(1, last + 1):
            ws.cell(row=row, column=col).border = Border(bottom=_thin)
        row += 1
    ws.freeze_panes = "A3"
    _xl_print_setup(ws, meta)
    ws.oddFooter.left.text = f"{TITLE} · {meta['range_label']}"
    return ws, row


def _xl_monthly(wb, data, meta):
    ws = wb.create_sheet("รายเดือน")
    names = [v["name"] for v in data["vehicles"]]
    headers = ["เดือน", *names, "รวม (บาท)", "ลิตร", "ราคาดีเซลเฉลี่ย", "ผลจากราคา", *[f"ลิตร {n}" for n in names]]
    last = len(headers)
    _xl_title(ws, meta, f"{TITLE} — รายเดือน", last)
    _xl_header(ws, 4, headers, "l" + "r" * (last - 1))
    for i, w in enumerate([12] + [12] * len(names) + [14, 10, 14, 13] + [11] * len(names), start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    r = 4
    for r, m in enumerate(data["months"], start=5):
        _put(ws, r, 1, m["label"], align=Alignment(indent=1))
        for j, n in enumerate(names, start=2):
            _put(ws, r, j, m["by_vehicle"].get(n) or 0, _f(10, color=C_INK if m["by_vehicle"].get(n) else C_MUTED), fmt=MONEY)
        c = len(names) + 2
        _put(ws, r, c, m["amount"], _f(10, True), fmt=MONEY)
        _put(ws, r, c + 1, m["liters"], fmt="#,##0")
        _put(ws, r, c + 2, m["diesel_price"], fmt="0.00")
        eff = m["price_effect"]
        _put(ws, r, c + 3, eff, _f(10, False, C_UP if eff > 0.5 else C_DOWN if eff < -0.5 else C_MUTED), fmt="+#,##0;-#,##0;0")
        for j, n in enumerate(names, start=c + 4):
            _put(ws, r, j, m["liters_by_vehicle"].get(n) or 0, fmt="#,##0.0")
        for col in range(1, last + 1):
            ws.cell(row=r, column=col).border = Border(bottom=_thin)
    tr = r + 1
    _put(ws, tr, 1, "รวม", _f(10, True), _fill(C_ACCENT_SOFT), align=Alignment(indent=1))
    for col in range(2, last + 1):
        letter = get_column_letter(col)
        if headers[col - 1] == "ราคาดีเซลเฉลี่ย":
            _put(ws, tr, col, None, fill=_fill(C_ACCENT_SOFT))
            continue
        fmt = "#,##0" if headers[col - 1].startswith("ลิตร") else "+#,##0;-#,##0;0" if headers[col - 1] == "ผลจากราคา" else MONEY
        _put(ws, tr, col, f"=SUM({letter}5:{letter}{r})", _f(10, True), _fill(C_ACCENT_SOFT), fmt=fmt)
    for col in range(1, last + 1):
        ws.cell(row=tr, column=col).border = Border(top=_medium)
    ws.freeze_panes = "B5"
    _xl_print_setup(ws, meta, landscape_mode=True, repeat_row=4)
    ws.oddFooter.left.text = f"{TITLE} · {meta['range_label']}"
    return ws, 5, r


def _xl_chart(summary_ws, month_ws, data, first, last_row, anchor_row):
    n = len(data["vehicles"])
    bar = BarChart()
    bar.type, bar.grouping, bar.overlap = "col", "stacked", 100
    bar.title = "ค่าน้ำมันรายเดือนแยกคัน และราคาดีเซล"
    bar.add_data(Reference(month_ws, min_col=2, max_col=n + 1, min_row=4, max_row=last_row), titles_from_data=True)
    bar.set_categories(Reference(month_ws, min_col=1, min_row=first, max_row=last_row))
    for i, s in enumerate(bar.series):
        s.graphicalProperties.solidFill = _color(i)
        s.graphicalProperties.line.noFill = True
    bar.gapWidth = 40
    bar.y_axis.numFmt = "#,##0"
    bar.y_axis.scaling.min = 0
    bar.y_axis.title = "บาท"
    bar.x_axis.delete = bar.y_axis.delete = False
    line = LineChart()
    price_col = n + 4
    line.add_data(Reference(month_ws, min_col=price_col, min_row=4, max_row=last_row), titles_from_data=True)
    line.series[0].graphicalProperties.line.solidFill = C_INK
    line.series[0].graphicalProperties.line.width = 22000
    line.series[0].smooth = False
    line.y_axis.axId = 200
    line.y_axis.title = "บาท/ลิตร"
    line.y_axis.crosses = "max"
    line.y_axis.majorGridlines = None
    line.y_axis.delete = False
    bar += line
    bar.legend.position = "b"
    bar.title.overlay = False
    bar.width, bar.height = 26, 11
    summary_ws.add_chart(bar, f"A{anchor_row + 2}")

    liters = BarChart()
    liters.type, liters.grouping, liters.overlap = "col", "stacked", 100
    liters.title = "ปริมาณน้ำมันที่เติมรายเดือน (ลิตร)"
    first_l = n + 6
    liters.add_data(Reference(month_ws, min_col=first_l, max_col=first_l + n - 1, min_row=4, max_row=last_row),
                    titles_from_data=True)
    liters.set_categories(Reference(month_ws, min_col=1, min_row=first, max_row=last_row))
    for i, s in enumerate(liters.series):
        s.graphicalProperties.solidFill = _color(i)
        s.graphicalProperties.line.noFill = True
    liters.gapWidth = 40
    liters.y_axis.numFmt = "#,##0"
    liters.y_axis.scaling.min = 0
    liters.y_axis.title = "ลิตร"
    liters.x_axis.delete = liters.y_axis.delete = False
    liters.legend.position = "b"
    liters.title.overlay = False
    liters.width, liters.height = 26, 10
    summary_ws.add_chart(liters, f"A{anchor_row + 26}")


def _xl_issues(wb, data, meta):
    if not data["issues"]:
        return
    ws = wb.create_sheet("ต้องตรวจ")
    ws.sheet_properties.tabColor = C_WARN
    headers = ["วันที่", "รถ", "ปัญหา", "ยอด", "ปั๊ม", "หมายเหตุ"]
    _xl_title(ws, meta, f"{TITLE} — รายการที่ต้องตรวจ", len(headers))
    _xl_header(ws, 4, headers, "lllrll")
    for i, w in enumerate([12, 10, 32, 12, 30, 60], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, f in enumerate(reversed(data["issues"]), start=5):
        _put(ws, r, 1, _d(f["date"]), fmt="dd/mm/yyyy", align=Alignment(horizontal="left", vertical="top"))
        _put(ws, r, 2, f["vehicle"], align=Alignment(vertical="top"))
        _put(ws, r, 3, f["issue"], _f(10, True, C_WARN), align=Alignment(vertical="top", wrap_text=True))
        _put(ws, r, 4, f["amount"], fmt=MONEY, align=Alignment(vertical="top"))
        _put(ws, r, 5, f["station"], align=Alignment(vertical="top", wrap_text=True))
        _put(ws, r, 6, f["note"], _f(9, color=C_MUTED), align=Alignment(vertical="top", wrap_text=True))
        for col in range(1, len(headers) + 1):
            ws.cell(row=r, column=col).border = Border(bottom=_thin)
    ws.freeze_panes = "A5"
    _xl_print_setup(ws, meta, landscape_mode=True, repeat_row=4)
    ws.oddFooter.left.text = f"{TITLE} · {meta['range_label']}"


def _xl_fills(wb, data, meta):
    from openpyxl.worksheet.table import Table as XLTable, TableStyleInfo

    ws = wb.create_sheet("ประวัติการเติม")
    cols = [("วันที่", 12), ("รถ", 9), ("ทะเบียน", 10), ("น้ำมัน", 9), ("ผลิตภัณฑ์", 18), ("บาท/ลิตร", 10),
            ("ลิตร", 10), ("ยอด", 12), ("เลขไมล์", 11), ("ปั๊ม", 28), ("ประมาณการ", 11), ("ปัญหา", 28), ("หมายเหตุ", 50)]
    _xl_title(ws, meta, f"{TITLE} — ประวัติการเติม", len(cols))
    _xl_header(ws, 4, [c[0] for c in cols], "lllllrrrrllll")
    for i, (_, w) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    from vehicle_costs import FUEL_LABEL, issue_of
    r = 4
    for r, f in enumerate(data["fills"], start=5):
        vals = [_d(f["date"]), f["vehicle"], f["plate"], FUEL_LABEL.get(f["fuel_type"], f["fuel_type"]), f["product"],
                f["price"], f["liters"], f["amount"], f["odometer"], f["station"],
                "ใช่" if f["estimated"] else "", issue_of(f), f["note"]]
        fmts = ["dd/mm/yyyy", None, None, None, None, "0.00", "#,##0.000", MONEY, "#,##0", None, None, None, None]
        for c, (v, fm) in enumerate(zip(vals, fmts), start=1):
            _put(ws, r, c, v, fmt=fm, align=Alignment(vertical="top", horizontal="left" if c == 1 else None))
    tbl = XLTable(displayName="tbl_fills", ref=f"A4:{get_column_letter(len(cols))}{r}")
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(tbl)
    ws.freeze_panes = "B5"
    _xl_print_setup(ws, meta, landscape_mode=True, repeat_row=4)
    ws.oddFooter.left.text = f"{TITLE} · {meta['range_label']}"


RECON_COLORS = {"matched": C_DOWN, "slip_only": C_UP, "receipt_only": C_WARN, "vehicle_mismatch": C_WARN,
                "before_ledger": C_MUTED}


def _recon_rows(data):
    """Timeline rows from the period the bank ledger covers (older receipts have no slips to match)."""
    return [r for r in data["timeline"]["rows"] if r["status"] != "before_ledger"]


def _xl_recon(wb, data, meta):
    tl = data.get("timeline")
    if not tl:
        return
    ws = wb.create_sheet("สลิป vs ใบเสร็จ")
    cols = [("วันที่", 12), ("รถ", 10), ("ยอด", 12), ("สถานะ", 30), ("ใบเสร็จ: ปั๊ม", 28), ("ใบเสร็จ: ลิตร", 12),
            ("สลิป: วันที่", 12), ("สลิป: ผู้รับ", 30), ("สลิป: บันทึก", 34)]
    _xl_title(ws, meta, f"{TITLE} — สลิปโอน เทียบ ใบเสร็จ", len(cols))
    c = tl["counts"]
    _put(ws, 3, 1, (f"ตั้งแต่ {tl['ledger_start']}: ตรงกัน {c['matched']} · มีสลิปไม่มีใบกำกับภาษี {c['slip_only']} "
                    f"({tl['amounts']['slip_only']:,.0f} บาท) · มีใบเสร็จไม่พบสลิป {c['receipt_only']} "
                    f"({tl['amounts']['receipt_only']:,.0f} บาท) · ลงชื่อรถไม่ตรง {c['vehicle_mismatch']}"),
         _f(9, color=C_MUTED, italic=True))
    _xl_header(ws, 4, [x[0] for x in cols], "llrllrlll")
    for i, (_, w) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, row in enumerate(_recon_rows(data), start=5):
        f = row["fill"] if row["fill"] and not row["fill"]["estimated"] else None
        sl = row["slip"]
        vals = [_d(row["date"]), row["vehicle"], row["amount"], row["status_label"],
                f["station"] if f else "", f["liters"] if f else None,
                _d(sl["date"]) if sl else None, sl["receiver"] if sl else "", sl["memo"] if sl else ""]
        fmts = ["dd/mm/yyyy", None, MONEY, None, None, "#,##0.0", "dd/mm/yyyy", None, None]
        for col, (v, fm) in enumerate(zip(vals, fmts), start=1):
            font = _f(10, True, RECON_COLORS[row["status"]]) if col == 4 else None
            _put(ws, r, col, v, font, fmt=fm, align=Alignment(vertical="top", wrap_text=col in (5, 8, 9),
                                                              horizontal="left" if col in (1, 7) else None))
            ws.cell(row=r, column=col).border = Border(bottom=_thin)
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{max(5, 4 + len(_recon_rows(data)))}"
    ws.freeze_panes = "A5"
    _xl_print_setup(ws, meta, landscape_mode=True, repeat_row=4)
    ws.oddFooter.left.text = f"{TITLE} · {meta['range_label']}"


def build_vehicle_xlsx(data, output):
    meta = _meta(data)
    wb = Workbook()
    summary, end_row = _xl_summary(wb, data, meta)
    month_ws, first, last_row = _xl_monthly(wb, data, meta)
    _xl_chart(summary, month_ws, data, first, last_row, end_row)
    _xl_recon(wb, data, meta)
    _xl_issues(wb, data, meta)
    _xl_fills(wb, data, meta)
    wb.active = 0
    wb.properties.title = TITLE
    wb.save(output)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _pdf_vehicle_table(data):
    best = min((v for v in data["vehicles"] if v["baht_per_km"]), key=lambda v: v["baht_per_km"], default=None)
    widths = [CONTENT_W - 345, 45, 50, 60, 55, 75, 60]
    rows = [[P("รถ", "head"), P("น้ำมัน", "head"), P("กม./ลิตร", "head_r"), P("วิ่ง/เดือน", "head_r"),
             P("บาท/กม.", "head_r"), P("ค่าน้ำมันรวม", "head_r"), P("ต้องตรวจ", "head_r")]]
    style = _base_table_style()
    for i, v in enumerate(data["vehicles"]):
        name = (f'<font color="#{_color(i)}">■</font> <b>{escape(v["name"])}</b> '
                f'<font color="#{C_MUTED}" size="7.5">{escape(v["plate"])}</font>')
        bpk = f"{v['baht_per_km']:.2f}" if v["baht_per_km"] else "–"
        if best and v is best:
            bpk = f'<b>{bpk}</b><br/><font color="#{C_DOWN}" size="7">ถูกสุด</font>'
        rows.append([
            Paragraph(name, ST["cell"]), P(v["fuel_label"], "cell_muted"),
            P(f"{v['km_per_l']:.1f}" if v["km_per_l"] else "–", "cell_r"),
            P(f"{v['km_per_month']:,.0f} กม." if v["km_per_month"] else "–", "cell_r"),
            Paragraph(bpk, ST["cell_r"]), P(fmt_baht(v["amount"]), "cell_r"),
            P(v["issues"] or "–", "cell_r", bold=bool(v["issues"]), color=C_WARN if v["issues"] else C_MUTED),
        ])
        if v["efficiency_note"]:
            rows.append(["", PW(v["efficiency_note"], sum(widths[1:]), "cell_muted"), "", "", "", "", ""])
            r = len(rows) - 1
            style += [("SPAN", (1, r), (-1, r)), ("TEXTCOLOR", (1, r), (-1, r), _hex(C_WARN)),
                      ("TOPPADDING", (0, r), (-1, r), 0)]
    t = Table(rows, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle(style))
    return t


def _pdf_month_chart(data, key="by_vehicle"):
    months, names = data["months"], [v["name"] for v in data["vehicles"]]
    height = 190
    d = Drawing(CONTENT_W, height)
    chart = VerticalBarChart()
    chart.x, chart.y = 45, 40
    chart.width, chart.height = CONTENT_W - 90, height - 55
    chart.data = [[m[key].get(n, 0) for m in months] for n in names]
    chart.categoryAxis.style = "stacked"
    vmax, step = _nice_max(max(sum(m[key].values()) for m in months))
    chart.valueAxis.valueMin, chart.valueAxis.valueMax, chart.valueAxis.valueStep = 0, vmax, step
    chart.valueAxis.labelTextFormat = lambda v: f"{v:,.0f}"
    for ax in (chart.valueAxis.labels, chart.categoryAxis.labels):
        ax.fontName, ax.fontSize, ax.fillColor = FONT, 6.5, _hex(C_MUTED)
    chart.valueAxis.strokeColor = None
    chart.valueAxis.visibleGrid = 1
    chart.valueAxis.gridStrokeColor = _hex("E5E7EB")
    chart.categoryAxis.categoryNames = [m["label"] for m in months]
    chart.categoryAxis.labels.angle = 45
    chart.categoryAxis.labels.boxAnchor = "ne"
    chart.categoryAxis.tickDown = 0
    for i in range(len(names)):
        chart.bars[i].fillColor = _hex(_color(i))
        chart.bars[i].strokeColor = None
    chart.barSpacing, chart.groupSpacing = 0, 4
    d.add(chart)

    # Diesel price on its own right-hand scale, drawn over the bars.
    prices = [m["diesel_price"] for m in months]
    known = [p for p in prices if p]
    if known:
        pmin, pmax = 25, max(45, max(known) + 2)
        line = HorizontalLineChart()
        line.x, line.y, line.width, line.height = chart.x, chart.y, chart.width, chart.height
        line.data = [[p if p else None for p in prices]]
        line.valueAxis.valueMin, line.valueAxis.valueMax = pmin, pmax
        line.valueAxis.visible = 0
        line.categoryAxis.visible = 0
        line.lines[0].strokeColor = _hex(C_INK)
        line.lines[0].strokeWidth = 1.4
        d.add(line)
        for val in range(30, int(pmax) + 1, 5):
            y = chart.y + (val - pmin) / (pmax - pmin) * chart.height
            d.add(String(chart.x + chart.width + 6, y - 2, f"{val}", fontName=FONT, fontSize=6.5, fillColor=_hex(C_MUTED)))
        d.add(String(chart.x + chart.width + 6, chart.y + chart.height + 6, "บาท/ลิตร", fontName=FONT, fontSize=6.5,
                     fillColor=_hex(C_MUTED)))
    return d


def _pdf_legend(data):
    cells = []
    for i, v in enumerate(data["vehicles"]):
        sw = Drawing(9, 9)
        sw.add(Rect(0, 0, 9, 9, fillColor=_hex(_color(i)), strokeColor=None))
        cells += [sw, P(v["name"], "cell_muted")]
    line = Drawing(14, 9)
    line.add(Rect(0, 3.5, 14, 1.6, fillColor=_hex(C_INK), strokeColor=None))
    cells += [line, P("ราคาดีเซลเฉลี่ย", "cell_muted")]
    t = Table([cells], colWidths=[14, 60] * (len(cells) // 2))
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 2),
                           ("RIGHTPADDING", (0, 0), (-1, -1), 2)]))
    return t


def _pdf_month_table(data):
    widths = [70, 45, 70, 95, 90, CONTENT_W - 370]
    rows = [[P("เดือน", "head"), P("เติม", "head_r"), P("ลิตร", "head_r"), P("ค่าน้ำมัน", "head_r"),
             P("ราคาดีเซลเฉลี่ย", "head_r"), P("ผลจากราคา", "head_r")]]
    for m in reversed(data["months"]):
        eff = m["price_effect"]
        rows.append([P(m["label"]), P(m["fills"], "cell_r"), P(f"{m['liters']:,.0f}", "cell_r"),
                     P(fmt_baht(m["amount"]), "cell_r"),
                     P(f"{m['diesel_price']:.2f}" if m["diesel_price"] else "–", "cell_r"),
                     P(f"{eff:+,.0f}", "cell_r", bold=abs(eff) > 1000,
                       color=C_UP if eff > 0.5 else C_DOWN if eff < -0.5 else C_MUTED)])
    total_eff = sum(m["price_effect"] for m in data["months"])
    rows.append([P("รวม", bold=True), P(data["kpi"]["fills"], "cell_r", bold=True),
                 P(f"{data['kpi']['liters']:,.0f}", "cell_r", bold=True), P(fmt_baht(data["kpi"]["amount"]), "cell_r", bold=True),
                 P(""), P(f"{total_eff:+,.0f}", "cell_r", bold=True, color=C_UP if total_eff > 0 else C_DOWN)])
    t = Table(rows, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle(_base_table_style() + [
        ("BACKGROUND", (0, -1), (-1, -1), _hex(C_ACCENT_SOFT)), ("LINEABOVE", (0, -1), (-1, -1), 1, _hex(C_INK)),
    ]))
    return t


def _pdf_issues(data):
    widths = [62, 48, 150, 65, CONTENT_W - 325]
    rows = [[P("วันที่", "head"), P("รถ", "head"), P("ปัญหา", "head"), P("ยอด", "head_r"), P("ปั๊ม", "head")]]
    for f in reversed(data["issues"]):
        rows.append([P(fmt_date(_d(f["date"])), "cell_muted"), P(f["vehicle"]), PW(f["issue"], widths[2], bold=True),
                     P(fmt_baht(f["amount"]), "cell_r"), PW(f["station"], widths[4], "cell_muted")])
    t = Table(rows, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle(_base_table_style() + [("TEXTCOLOR", (2, 1), (2, -1), _hex(C_WARN))]))
    return t


def _pdf_fills(data):
    widths = [58, 42, 46, 50, 62, 58, CONTENT_W - 316]
    rows = [[P("วันที่", "head"), P("รถ", "head"), P("บาท/ลิตร", "head_r"), P("ลิตร", "head_r"),
             P("ยอด", "head_r"), P("เลขไมล์", "head_r"), P("ปั๊ม", "head")]]
    style = _base_table_style()
    for f in reversed(data["fills"]):
        est = f["estimated"]
        rows.append([
            P(fmt_date(_d(f["date"])), "cell_muted"), P(f["vehicle"]),
            P(f"{f['price']:.2f}{'*' if est else ''}" if f["price"] else "–", "cell_r", color=C_MUTED if est else None),
            P(f"{'~' if est else ''}{f['liters']:,.1f}" if f["liters"] else "–", "cell_r", color=C_MUTED if est else None),
            P(fmt_baht(f["amount"]), "cell_r"),
            P(f"{f['odometer']:,.0f}" if f["odometer"] else "–", "cell_r", color=None if f["odometer"] else C_MUTED),
            PW(f["station"] + (" · ไม่มีใบกำกับภาษี" if est else ""), widths[6], "cell_muted", max_lines=2),
        ])
    t = Table(rows, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle(style + [("ROWBACKGROUNDS", (0, 1), (-1, -1), ["white", _hex(C_SURFACE_ALT)])]))
    return t


def _pdf_recon(data):
    widths = [58, 42, 58, 115, 118, CONTENT_W - 391]
    rows = [[P("วันที่", "head"), P("รถ", "head"), P("ยอด", "head_r"), P("ใบเสร็จ", "head"), P("สลิปโอน", "head"),
             P("สถานะ", "head")]]
    for r in _recon_rows(data):
        if r["status"] == "matched":
            continue
        f = r["fill"] if r["fill"] and not r["fill"]["estimated"] else None
        sl = r["slip"]
        rows.append([
            P(fmt_date(_d(r["date"])), "cell_muted"), P(r["vehicle"]), P(fmt_baht(r["amount"]), "cell_r"),
            PW(f"{f['station']} · {f['liters']:.1f} ล." if f else "—", widths[3], "cell_muted", max_lines=2),
            PW(f"{sl['receiver']} · {sl['memo']}" if sl else "—", widths[4], "cell_muted", max_lines=2),
            Paragraph(f'<font color="#{RECON_COLORS[r["status"]]}"><b>{escape(r["status_label"])}</b></font>', ST["cell"]),
        ])
    t = Table(rows, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle(_base_table_style()))
    return t


def usage_headline(u):
    t = u["total_complete"]
    return (f"{', '.join(u['complete_names'])}: {t['verdict']} ({t['liters_change']:+.1f}%) — "
            f"{t['before']['liters']:,.0f} → {t['after']['liters']:,.0f} ลิตร/เดือน · ค่าน้ำมัน/เดือน "
            f"{t['spend_change']:+,.0f} บาท = จากราคา {t['price_effect']:+,.0f} + จากปริมาณ {t['volume_effect']:+,.0f}")


def _pdf_usage(data):
    u = data["usage"]
    widths = [CONTENT_W - 395, 55, 50, 50, 90, 60, 45, 45]
    rows = [[P("รถ", "head"), P("ลิตร/ด. ก่อน", "head_r"), P("หลัง", "head_r"), P("เปลี่ยน", "head_r"), P("ผล", "head"),
             P("บาท/ด. เปลี่ยน", "head_r"), P("ปริมาณ", "head_r"), P("ราคา", "head_r")]]
    for r in u["vehicles"]:
        ch = r["liters_change"]
        col = C_UP if ch and ch >= 5 else C_DOWN if ch and ch <= -5 else None
        rows.append([
            P(r["name"], bold=True), P(f"{r['before']['liters']:,.0f}", "cell_r"), P(f"{r['after']['liters']:,.0f}", "cell_r"),
            P(f"{ch:+.1f}%" if ch is not None else "–", "cell_r", color=col),
            PW("ใบเสร็จไม่ครบ ตีความไม่ได้" if r["missing_liters"] else r["verdict"], widths[4],
               "cell_muted" if r["missing_liters"] else "cell"),
            P(f"{r['spend_change']:+,.0f}", "cell_r"),
            P(f"{r['volume_effect']:+,.0f}", "cell_r", color=C_UP if r["volume_effect"] > 0.5 else C_DOWN),
            P(f"{r['price_effect']:+,.0f}", "cell_r", color=C_UP if r["price_effect"] > 0.5 else C_DOWN),
        ])
    t = Table(rows, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle(_base_table_style()))
    return [CondPageBreak(7 * 28), Paragraph("ราคาขึ้นแล้ว เราใช้น้ำมันเปลี่ยนไหม", ST["h2"]),
            Paragraph(f"เฉลี่ยต่อเดือน ก่อน ({escape(u['before_label'])}) เทียบ หลังราคาขึ้น ({escape(u['after_label'])}) "
                      "· จากปริมาณ = ลิตรที่เปลี่ยน × ราคาเดิม, จากราคา = ลิตรใหม่ × ราคาที่เปลี่ยน", ST["muted"]),
            Spacer(1, 4), _callout(f"<b>{escape(usage_headline(u))}</b>", fg=C_ACCENT, bg=C_ACCENT_SOFT), Spacer(1, 6), t]


def build_vehicle_pdf(data, output):
    meta = _meta(data)
    k = data["kpi"]
    gen = meta["generated_at"]
    d_now, d_base = data["latest_price"].get("diesel"), data["baseline"].get("diesel")
    story = [
        header_band(TITLE, "ทุกคัน", "",
                    [f"ข้อมูล {escape(meta['range_label'])}", f"{k['fills']} ครั้ง · {len(data['vehicles'])} คัน",
                     f"จัดทำเมื่อ {fmt_date(gen.date())} {gen:%H:%M} น."]),
        Spacer(1, 12),
    ]
    diesel_sub = ""
    if d_now and d_base:
        pct = (d_now - d_base) / d_base * 100
        diesel_sub = (f'<font color="#{C_UP if pct > 0 else C_DOWN}"><b>{pct:+.1f}%</b></font> '
                      f"เทียบก่อน {escape(data['baseline_label'])} ({d_base:.2f})")
    story += [kpi_cards([
        ("ค่าน้ำมันรวม (บาท)", fmt_baht(k["amount"]), f"{k['liters']:,.0f} ลิตร"),
        ("เฉลี่ย/เดือน 3 เดือนล่าสุด", fmt_baht(k["monthly_avg_recent"]), f"~{k['monthly_liters_recent']:,.0f} ลิตร/เดือน"),
        ("จ่ายเพิ่มเพราะราคาน้ำมัน", f'<font color="#{C_UP}">+{fmt_baht(k["price_effect_since"])}</font>',
         f"ตั้งแต่ {escape(data['baseline_label'])}"),
        ("ราคาดีเซลล่าสุด (บาท/ลิตร)", f"{d_now:.2f}" if d_now else "–", diesel_sub),
    ]), Spacer(1, 10)]
    for v in data["vehicles"]:
        if v["missing_liters"]:
            story += [_callout(f"<b>{escape(v['name'])}</b>: ใบเสร็จขาดประมาณ <b>{v['missing_liters']:,.0f} ลิตร</b> "
                               f"(~{v['missing_liters'] * (d_now or 0):,.0f} บาท ที่ราคาปัจจุบัน) "
                               "เมื่อเทียบระยะทางที่วิ่งจริงกับอัตราสิ้นเปลืองปกติของคันนี้"), Spacer(1, 6)]
    story += [Paragraph("เทียบรายคัน", ST["h2"]),
              Paragraph("บาท/กม. คิดจากราคาน้ำมันล่าสุด · กม./ลิตร คิดจากเลขไมล์และลิตรที่เติมระหว่างนั้น", ST["muted"]),
              Spacer(1, 4), _pdf_vehicle_table(data)]
    story += [KeepTogether([Paragraph("ค่าน้ำมันรายเดือน (บาท)", ST["h2"]), _pdf_month_chart(data), _pdf_legend(data)])]
    story += [KeepTogether([Paragraph("ปริมาณน้ำมันที่เติมรายเดือน (ลิตร)", ST["h2"]),
                            _pdf_month_chart(data, "liters_by_vehicle"), _pdf_legend(data)])]
    story += _pdf_usage(data)
    story += [CondPageBreak(6 * 28), Paragraph("ราคาน้ำมันกระทบค่าใช้จ่ายเท่าไหร่", ST["h2"]),
              Paragraph(f"ผลจากราคา = ลิตรที่เติม × (ราคาที่จ่าย − ราคาเฉลี่ยก่อน {escape(data['baseline_label'])}) "
                        "· แดง = จ่ายแพงขึ้น, เขียว = ถูกลง", ST["muted"]),
              Spacer(1, 4), _pdf_month_table(data)]
    tl = data.get("timeline")
    if tl:
        c, a = tl["counts"], tl["amounts"]
        story += [CondPageBreak(6 * 28), Paragraph("สลิปโอน เทียบ ใบเสร็จ", ST["h2"]),
                  Paragraph(f"ช่วงที่มีสลิปในบัญชี (ตั้งแต่ {fmt_date(_d(tl['ledger_start']))}) · ตรงกัน {c['matched']} รายการ · "
                            f"มีสลิปไม่มีใบกำกับภาษี {c['slip_only']} ({fmt_baht(a['slip_only'])} บาท) · "
                            f"มีใบเสร็จไม่พบสลิป {c['receipt_only']} ({fmt_baht(a['receipt_only'])} บาท) · "
                            f"ลงชื่อรถไม่ตรง {c['vehicle_mismatch']}", ST["muted"]),
                  Spacer(1, 4), _pdf_recon(data)]
    if data["issues"]:
        story += [CondPageBreak(6 * 28), Paragraph("รายการที่ต้องตรวจ", ST["h2"]),
                  Paragraph(f"{len(data['issues'])} รายการ", ST["muted"]), Spacer(1, 4), _pdf_issues(data)]
    story += [PageBreak(), Paragraph("ประวัติการเติมทั้งหมด", ST["h2"]),
              Paragraph("* / ~ = ไม่มีใบกำกับภาษี ประมาณลิตรจากราคาใบเสร็จใกล้วันที่สุด", ST["muted"]),
              Spacer(1, 4), _pdf_fills(data)]
    footer = f"{TITLE} · {meta['range_label']} · จัดทำ {fmt_date(gen.date())}"
    SimpleDocTemplate(
        output, pagesize=A4, title=TITLE, author="Ledger",
        topMargin=MARGIN, bottomMargin=MARGIN + 0.3 * 28.35, leftMargin=MARGIN, rightMargin=MARGIN,
    ).build(story, canvasmaker=lambda *a, **kw: _NumberedCanvas(*a, footer_left=footer, **kw))
